from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Iterable

from db import db_conn

RAW_TABLE_BY_SOURCE = {
    "movement": "raw_movement",
    "price_list": "raw_price_list",
    "sales_by_price_type": "raw_sales_by_price_type",
    "on_hand": "raw_on_hand",
    "frozen_inventory": "raw_frozen_inventory",
    "tx_seafood_abc_price": "raw_tx_seafood_abc_price",
    "live_fresh_suggestion": "raw_live_fresh_suggestion",
    "article_manual": "raw_article_manual",
    "monthly_pi_form": "raw_monthly_pi_form",
    "unclassified": "raw_unclassified",
}


@dataclass
class ParsedSheet:
    index: int
    name: str
    visibility: str
    rows: list[list[Any]]
    header_row_number: int | None = None
    headers: list[str] = field(default_factory=list)
    coverage_type: str = "UNKNOWN"
    coverage_start: date | None = None
    coverage_end: date | None = None
    date_semantics_status: str = "UNRESOLVED"
    date_evidence: dict[str, Any] = field(default_factory=dict)
    row_dates: dict[int, date] = field(default_factory=dict)


@dataclass
class ParsedFile:
    filename: str
    sha256: str
    size_bytes: int
    detected_format: str
    source_type: str
    source_type_confidence: float
    sheets: list[ParsedSheet]
    coverage_type: str = "UNKNOWN"
    coverage_start: date | None = None
    coverage_end: date | None = None
    header_period_start: date | None = None
    header_period_end: date | None = None
    header_period_text: str | None = None
    filename_date_candidate: date | None = None
    filename_time_candidate: str | None = None
    report_generated_at: datetime | None = None
    date_semantics_status: str = "UNRESOLVED"
    date_evidence: dict[str, Any] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


def _jsonable(v: Any) -> Any:
    if v is None or isinstance(v, (str, int, float, bool)):
        return v
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    return str(v)


def _cell_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()



def _trim_trailing_blank_rows(rows: list[list[Any]]) -> list[list[Any]]:
    end = len(rows)
    while end > 0 and all(not _cell_text(v) for v in rows[end - 1]):
        end -= 1
    return rows[:end]

def _excel_col(n: int) -> str:
    out = ""
    while n:
        n, r = divmod(n - 1, 26)
        out = chr(65 + r) + out
    return out


def detect_format(data: bytes, filename: str) -> str:
    if data.startswith(b"PK\x03\x04") and zipfile.is_zipfile(io.BytesIO(data)):
        return "XLSX_ZIP_XML"
    if data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        return "LEGACY_XLS_OLE"
    # Older Excel BIFF workbook stream without an OLE compound-file wrapper.
    if len(data) >= 2 and data[:2] in {b'\x09\x00', b'\x09\x02', b'\x09\x04', b'\x09\x08'}:
        return "LEGACY_XLS_BIFF"
    if data.startswith(b"\xff\xfe") or data.startswith(b"\xfe\xff"):
        return "UTF16_TEXT_EXPORT"
    if data.startswith(b"\xef\xbb\xbf"):
        return "UTF8_TEXT_EXPORT"
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext in {"csv", "txt"}:
        return "TEXT_EXPORT"
    return "UNKNOWN"


def classify_source(filename: str) -> tuple[str, float]:
    n = filename.lower()
    rules = [
        ("article manual", "article_manual"),
        ("monthly pi", "monthly_pi_form"),
        ("abc price", "tx_seafood_abc_price"),
        ("live & fresh", "live_fresh_suggestion"),
        ("f inv", "frozen_inventory"),
        ("frozen", "frozen_inventory"),
        ("sales", "sales_by_price_type"),
        ("price", "price_list"),
        ("mvt", "movement"),
        ("movement", "movement"),
        ("on hand", "on_hand"),
    ]
    for needle, source in rules:
        if needle in n:
            return source, 0.95

    # OH reports are frequently named like:
    #   oh-092826.xls, OH_092826.xls, OH 092826.xls
    # Match OH as a standalone filename token, not as part of another word.
    if re.search(r"(^|[^a-z0-9])oh([^a-z0-9]|$)", n):
        return "on_hand", 0.95

    return "unclassified", 0.0


def _filename_datetime_candidates(filename: str) -> tuple[date | None, str | None]:
    # Supports MMDDYY[-_]HHMM and YYYYMMDD patterns as candidate metadata only.
    n = filename
    m = re.search(r"(?<!\d)(20\d{2})(\d{2})(\d{2})(?!\d)", n)
    if m:
        try:
            d = date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
            return d, None
        except ValueError:
            pass
    m = re.search(r"(?<!\d)(\d{2})(\d{2})(\d{2})(?:[-_](\d{2})(\d{2}))?(?!\d)", n)
    if m:
        try:
            yy = int(m.group(3))
            d = date(2000 + yy, int(m.group(1)), int(m.group(2)))
            t = f"{m.group(4)}:{m.group(5)}:00" if m.group(4) else None
            return d, t
        except ValueError:
            pass
    return None, None


def _read_xlsx(data: bytes) -> list[ParsedSheet]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError('openpyxl is not installed. Run INSTALL_ONCE.bat.') from exc
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=False, data_only=False)
    result: list[ParsedSheet] = []
    for idx, ws in enumerate(wb.worksheets, 1):
        visibility = str(ws.sheet_state or "visible").upper().replace("VERYHIDDEN", "VERY_HIDDEN")
        rows: list[list[Any]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append([_jsonable(v) for v in row])
        rows = _trim_trailing_blank_rows(rows)
        result.append(ParsedSheet(idx, ws.title, visibility, rows))
    return result


def _read_xls(data: bytes) -> list[ParsedSheet]:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError('xlrd is not installed. Run INSTALL_ONCE.bat.') from exc
    book = xlrd.open_workbook(file_contents=data, on_demand=True)
    result: list[ParsedSheet] = []
    for idx, name in enumerate(book.sheet_names(), 1):
        sh = book.sheet_by_name(name)
        rows: list[list[Any]] = []
        for r in range(sh.nrows):
            vals: list[Any] = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        vals.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode).isoformat(sep=" "))
                    except Exception:
                        vals.append(str(cell.value))
                else:
                    vals.append(_jsonable(cell.value))
            rows.append(vals)
        visibility = "VISIBLE"
        try:
            vis = book.sheet_visibility[idx - 1]
            visibility = {0: "VISIBLE", 1: "HIDDEN", 2: "VERY_HIDDEN"}.get(vis, "VISIBLE")
        except Exception:
            pass
        rows = _trim_trailing_blank_rows(rows)
        result.append(ParsedSheet(idx, name, visibility, rows))
    book.release_resources()
    return result


def _read_text(data: bytes, detected_format: str) -> list[ParsedSheet]:
    encodings = ["utf-16", "utf-8-sig", "utf-8", "cp1252"]
    text = None
    for enc in encodings:
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Text export could not be decoded safely")
    lines = text.splitlines()
    delimiter = "\t" if sum(line.count("\t") for line in lines[:20]) >= sum(line.count(",") for line in lines[:20]) else ","
    rows = _trim_trailing_blank_rows([line.split(delimiter) for line in lines])
    return [ParsedSheet(1, "__FILE__", "VISIBLE", rows)]


def _detect_header(sheet: ParsedSheet) -> None:
    best: tuple[int, int] | None = None
    for i, row in enumerate(sheet.rows[:60], 1):
        nonempty = sum(1 for v in row if _cell_text(v))
        if nonempty < 2:
            continue
        textual = sum(1 for v in row if _cell_text(v) and not re.fullmatch(r"[-+]?\d+(?:\.\d+)?", _cell_text(v)))
        score = nonempty + textual
        if best is None or score > best[1]:
            best = (i, score)
    if best:
        sheet.header_row_number = best[0]
        sheet.headers = [_cell_text(v) for v in sheet.rows[best[0] - 1]]


def _try_date(v: Any) -> date | None:
    if isinstance(v, date) and not isinstance(v, datetime):
        return v
    if isinstance(v, datetime):
        return v.date()
    s = _cell_text(v)
    if not s:
        return None
    s = s.replace(".", "/").replace("-", "/")
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            continue
    return None


def _analyze_sales_period(parsed: ParsedFile) -> None:
    joined = "\n".join(
        "\t".join(_cell_text(v) for v in row)
        for sh in parsed.sheets
        for row in sh.rows[:80]
    )
    # Explicit header From / To dates; filename is never promoted here.
    m = re.search(
        r"(?is)\bfrom\b\s*[:\-]?\s*(\d{1,2}[/-]\d{1,2}[/-]\d{2,4}).{0,120}?\bto\b\s*[:\-]?\s*(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})",
        joined,
    )
    if not m:
        # This report family also uses `Date: MM/DD/YYYY to MM/DD/YYYY`, which is an explicit period header.
        m = re.search(
            r"(?is)\bdate\s*:\s*(\d{1,2}[/-]\d{1,2}[/-]\d{2,4}).{0,40}?\bto\b\s*(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})",
            joined,
        )
    if not m:
        parsed.warnings.append("Sales report period not resolved from explicit From/To (or Date ... to ...) header; coverage remains UNKNOWN.")
        return
    d1, d2 = _try_date(m.group(1)), _try_date(m.group(2))
    if d1 and d2:
        parsed.header_period_start = min(d1, d2)
        parsed.header_period_end = max(d1, d2)
        parsed.header_period_text = m.group(0)[:255]
        parsed.coverage_start = parsed.header_period_start
        parsed.coverage_end = parsed.header_period_end
        parsed.coverage_type = "SINGLE_DATE" if d1 == d2 else "DATE_RANGE"
        parsed.date_semantics_status = "RESOLVED"
        parsed.date_evidence["sales_header_period"] = {"from": d1.isoformat(), "to": d2.isoformat()}


def _analyze_movement_dates(parsed: ParsedFile) -> None:
    all_dates: list[date] = []
    found_date_column = False
    for sh in parsed.sheets:
        _detect_header(sh)
        if not sh.headers or sh.header_row_number is None:
            continue
        date_col = None
        for idx, h in enumerate(sh.headers):
            key = re.sub(r"\s+", " ", h.strip().lower())
            if key in {"posting date", "pstng date", "post date", "business date"}:
                date_col = idx
                break
        if date_col is None:
            continue
        found_date_column = True
        dates: list[date] = []
        for source_row, row in enumerate(sh.rows[sh.header_row_number :], sh.header_row_number + 1):
            if date_col >= len(row):
                continue
            d = _try_date(row[date_col])
            if d:
                sh.row_dates[source_row] = d
                dates.append(d)
        if dates:
            sh.coverage_start, sh.coverage_end = min(dates), max(dates)
            sh.coverage_type = "SINGLE_DATE" if sh.coverage_start == sh.coverage_end else "MIXED_DATES"
            sh.date_semantics_status = "RESOLVED"
            sh.date_evidence["date_column"] = sh.headers[date_col]
            all_dates.extend(dates)
    if all_dates:
        parsed.coverage_start, parsed.coverage_end = min(all_dates), max(all_dates)
        parsed.coverage_type = "SINGLE_DATE" if parsed.coverage_start == parsed.coverage_end else "MIXED_DATES"
        parsed.date_semantics_status = "RESOLVED"
        parsed.date_evidence["movement_posting_dates"] = {"count": len(all_dates)}
    elif not found_date_column:
        parsed.warnings.append("Movement Posting Date column not explicitly resolved; coverage remains UNKNOWN.")


def parse_file(filename: str, data: bytes) -> ParsedFile:
    sha = hashlib.sha256(data).hexdigest()
    fmt = detect_format(data, filename)
    source_type, confidence = classify_source(filename)
    if fmt == "XLSX_ZIP_XML":
        sheets = _read_xlsx(data)
    elif fmt in {"LEGACY_XLS_OLE", "LEGACY_XLS_BIFF"}:
        sheets = _read_xls(data)
    elif fmt in {"UTF16_TEXT_EXPORT", "UTF8_TEXT_EXPORT", "TEXT_EXPORT"}:
        sheets = _read_text(data, fmt)
    else:
        # Try text as a safe fallback; if it fails, hold as unsupported rather than guessing.
        try:
            sheets = _read_text(data, fmt)
            fmt = "TEXT_EXPORT_FALLBACK"
        except Exception as exc:
            raise ValueError(f"Unsupported or unreadable file format for {filename}: {exc}") from exc

    d_candidate, t_candidate = _filename_datetime_candidates(filename)
    parsed = ParsedFile(
        filename=filename,
        sha256=sha,
        size_bytes=len(data),
        detected_format=fmt,
        source_type=source_type,
        source_type_confidence=confidence,
        sheets=sheets,
        filename_date_candidate=d_candidate,
        filename_time_candidate=t_candidate,
    )
    # Preserve an explicit report print/generated timestamp when present.
    for sh in parsed.sheets:
        for row in sh.rows[:120]:
            for i, cell in enumerate(row):
                if 'PRINTED' in _cell_text(cell).upper():
                    for nxt in row[i+1:i+4]:
                        txt = _cell_text(nxt)
                        for fmt_dt in ('%m/%d/%Y %I:%M:%S %p','%m/%d/%Y %H:%M:%S'):
                            try:
                                parsed.report_generated_at = datetime.strptime(txt, fmt_dt)
                                break
                            except ValueError:
                                pass
                        if parsed.report_generated_at:
                            break
                if parsed.report_generated_at:
                    break
            if parsed.report_generated_at:
                break
        if parsed.report_generated_at:
            break
    for sh in parsed.sheets:
        _detect_header(sh)
    if source_type == "sales_by_price_type":
        _analyze_sales_period(parsed)
    elif source_type == "movement":
        _analyze_movement_dates(parsed)
        if (
            parsed.coverage_start is not None
            and parsed.coverage_start == parsed.coverage_end
        ):
            ok, reason = movement_daily_close_filename_evidence(
                original_filename=parsed.filename,
                filename_date_candidate=parsed.filename_date_candidate,
                filename_time_candidate=parsed.filename_time_candidate,
                coverage_start_date=parsed.coverage_start,
                coverage_end_date=parsed.coverage_end,
                business_date=parsed.coverage_start,
            )
            if ok:
                parsed.date_evidence["movement_daily_close_filename_contract"] = {
                    "complete": True,
                    "reason": reason,
                    "business_date": parsed.coverage_start.isoformat(),
                    "upload_time_used": False,
                }
    elif source_type == "unclassified":
        parsed.warnings.append("Source type unresolved; rows will be stored in raw_unclassified and require QA classification.")
    return parsed


def preview_file(filename: str, data: bytes) -> dict[str, Any]:
    p = parse_file(filename, data)
    return {
        "filename": p.filename,
        "sha256": p.sha256,
        "bytes": p.size_bytes,
        "detected_format": p.detected_format,
        "source_type": p.source_type,
        "source_type_confidence": p.source_type_confidence,
        "sheet_count": len(p.sheets),
        "sheets": [
            {
                "index": s.index,
                "name": s.name,
                "visibility": s.visibility,
                "rows": len(s.rows),
                "columns": max((len(r) for r in s.rows), default=0),
                "header_row": s.header_row_number,
                "coverage_type": s.coverage_type,
                "coverage_start": s.coverage_start.isoformat() if s.coverage_start else None,
                "coverage_end": s.coverage_end.isoformat() if s.coverage_end else None,
            }
            for s in p.sheets
        ],
        "coverage_type": p.coverage_type,
        "coverage_start": p.coverage_start.isoformat() if p.coverage_start else None,
        "coverage_end": p.coverage_end.isoformat() if p.coverage_end else None,
        "header_period_start": p.header_period_start.isoformat() if p.header_period_start else None,
        "header_period_end": p.header_period_end.isoformat() if p.header_period_end else None,
        "filename_date_candidate": p.filename_date_candidate.isoformat() if p.filename_date_candidate else None,
        "filename_time_candidate": p.filename_time_candidate,
        "warnings": p.warnings,
    }


def _insert_batch(cur, label: str | None) -> int:
    import uuid
    cur.execute(
        "INSERT INTO ingest_batch (batch_key,batch_label,processing_status,file_count,source_row_count,error_count) VALUES (%s,%s,'PROCESSING',0,0,0)",
        (str(uuid.uuid4()), label or "Seafood upload"),
    )
    return int(cur.lastrowid)


def ingest_files(files: list[tuple[str, bytes]], batch_label: str | None = None) -> dict[str, Any]:
    parsed_files = [parse_file(name, data) for name, data in files]
    with db_conn() as conn:
        cur = conn.cursor(dictionary=True)
        try:
            batch_id = _insert_batch(cur, batch_label)
            cur.execute(
                """INSERT INTO pipeline_run (batch_id,pipeline_name,stage_name,run_number,processing_status,code_version)
                   VALUES (%s,'SEAFOOD_LOCAL_UPLOADER','RAW_INGEST',1,'STARTED','local-agent-v6')""",
                (batch_id,),
            )
            pipeline_run_id = int(cur.lastrowid)
            total_raw = 0
            total_rows_read = 0
            file_results: list[dict[str, Any]] = []
            for p in parsed_files:
                cur.execute("SELECT file_id FROM ingest_file WHERE file_sha256=%s ORDER BY file_id LIMIT 1", (p.sha256,))
                dup = cur.fetchone()
                duplicate_of = int(dup["file_id"]) if dup else None
                cur.execute(
                    """
                    INSERT INTO ingest_file (
                      batch_id,original_filename,declared_extension,detected_format,source_type,source_type_confidence,
                      file_size_bytes,file_sha256,duplicate_of_file_id,sheet_count,source_row_count,
                      coverage_type,coverage_start_date,coverage_end_date,header_period_start_date,header_period_end_date,
                      header_period_text,filename_date_candidate,filename_time_candidate,report_generated_at,date_semantics_status,date_evidence_json,
                      processing_status,error_count
                    ) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,0,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'PARSING',0)
                    """,
                    (
                        batch_id, p.filename, p.filename.rsplit('.',1)[-1].lower() if '.' in p.filename else None,
                        p.detected_format, p.source_type, p.source_type_confidence, p.size_bytes, p.sha256, duplicate_of,
                        len(p.sheets), p.coverage_type, p.coverage_start, p.coverage_end, p.header_period_start,
                        p.header_period_end, p.header_period_text, p.filename_date_candidate, p.filename_time_candidate, p.report_generated_at,
                        p.date_semantics_status, json.dumps(p.date_evidence),
                    ),
                )
                file_id = int(cur.lastrowid)
                raw_count = 0
                total_rows_read += sum(len(sh.rows) for sh in p.sheets)
                raw_table = RAW_TABLE_BY_SOURCE[p.source_type]
                for sh in p.sheets:
                    col_count = max((len(r) for r in sh.rows), default=0)
                    cur.execute(
                        """
                        INSERT INTO ingest_sheet (
                          file_id,sheet_index,sheet_name,sheet_visibility,row_count,column_count,header_row_number,
                          coverage_type,coverage_start_date,coverage_end_date,date_semantics_status,date_evidence_json,
                          processing_status,error_count
                        ) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'INGESTED',0)
                        """,
                        (
                            file_id, sh.index, sh.name, sh.visibility, len(sh.rows), col_count, sh.header_row_number,
                            sh.coverage_type, sh.coverage_start, sh.coverage_end, sh.date_semantics_status,
                            json.dumps(sh.date_evidence),
                        ),
                    )
                    sheet_id = int(cur.lastrowid)
                    header_vals = sh.headers or []
                    for c in range(1, col_count + 1):
                        header = header_vals[c-1] if c-1 < len(header_vals) else None
                        samples = []
                        for row in sh.rows[:25]:
                            if c - 1 < len(row) and _cell_text(row[c-1]):
                                samples.append(_jsonable(row[c-1]))
                            if len(samples) >= 5:
                                break
                        cur.execute(
                            """
                            INSERT INTO ingest_column (
                              sheet_id,column_position,source_column_ref,raw_header_text,normalized_header_candidate,
                              detected_data_type,type_confidence,sample_values_json
                            ) VALUES (%s,%s,%s,%s,%s,NULL,NULL,%s)
                            """,
                            (sheet_id, c, _excel_col(c), header, header.strip().lower() if header else None, json.dumps(samples)),
                        )
                    if duplicate_of is not None:
                        continue
                    for row_no, row in enumerate(sh.rows, 1):
                        payload = {_excel_col(i+1): _jsonable(v) for i, v in enumerate(row)}
                        raw_text = "\t".join(_cell_text(v) for v in row)
                        row_hash = hashlib.sha256(raw_text.encode("utf-8", errors="replace")).hexdigest()
                        observed_date = sh.row_dates.get(row_no)
                        observed_role = "POSTING_DATE" if observed_date and p.source_type == "movement" else None
                        sql = f"""
                            INSERT INTO {raw_table} (
                              batch_id,file_id,sheet_id,source_row_number,source_row_kind,observed_business_date,
                              observed_date_role,raw_payload,raw_text,row_sha256
                            ) VALUES (%s,%s,%s,%s,'UNKNOWN',%s,%s,%s,%s,%s)
                        """
                        cur.execute(
                            sql,
                            (batch_id, file_id, sheet_id, row_no, observed_date, observed_role,
                             json.dumps(payload, ensure_ascii=False), raw_text, row_hash),
                        )
                        raw_count += 1
                # Unknown source type is explicitly QA-held rather than silently classified.
                status = "QA_HOLD" if p.source_type == "unclassified" else "INGESTED"
                err_count = 1 if p.source_type == "unclassified" else 0
                cur.execute(
                    "UPDATE ingest_file SET source_row_count=%s,processing_status=%s,error_count=%s,completed_at=CURRENT_TIMESTAMP(6) WHERE file_id=%s",
                    (raw_count, status, err_count, file_id),
                )
                if p.source_type == "unclassified":
                    cur.execute(
                        """
                        INSERT INTO ingest_error (batch_id,file_id,pipeline_run_id,stage_name,severity,error_code,error_message,detail_json)
                        VALUES (%s,%s,%s,'SOURCE_CLASSIFICATION','ERROR','SOURCE_TYPE_UNRESOLVED',%s,%s)
                        """,
                        (batch_id, file_id, pipeline_run_id, f"Source type unresolved for {p.filename}", json.dumps({"warnings": p.warnings})),
                    )
                for warning in p.warnings:
                    cur.execute(
                        """INSERT INTO ingest_error (batch_id,file_id,pipeline_run_id,stage_name,severity,error_code,error_message,detail_json)
                           VALUES (%s,%s,%s,'DATE_OR_PARSE_SEMANTICS','WARN','SOURCE_WARNING',%s,%s)""",
                        (batch_id, file_id, pipeline_run_id, warning, json.dumps({"filename": p.filename})),
                    )
                total_raw += raw_count
                file_results.append({
                    "file_id": file_id, "filename": p.filename, "sha256": p.sha256,
                    "source_type": p.source_type, "raw_rows_inserted": raw_count,
                    "duplicate_of_file_id": duplicate_of, "status": status, "warnings": p.warnings,
                })
            cur.execute("SELECT COALESCE(SUM(error_count),0) AS ec FROM ingest_file WHERE batch_id=%s", (batch_id,))
            ec = int(cur.fetchone()["ec"])
            batch_status = "QA_HOLD" if ec else "INGESTED_PENDING_QA"
            cur.execute(
                "UPDATE ingest_batch SET file_count=%s,source_row_count=%s,error_count=%s,processing_status=%s,completed_at=CURRENT_TIMESTAMP(6) WHERE batch_id=%s",
                (len(parsed_files), total_raw, ec, batch_status, batch_id),
            )
            cur.execute(
                """UPDATE pipeline_run SET processing_status='COMPLETED',completed_at=CURRENT_TIMESTAMP(6),
                   rows_read=%s,rows_inserted=%s,rows_rejected=0,error_count=%s,detail_json=%s
                   WHERE pipeline_run_id=%s""",
                (total_rows_read, total_raw, ec, json.dumps({"file_count": len(parsed_files)}), pipeline_run_id),
            )
            conn.commit()
            return {"batch_id": batch_id, "status": batch_status, "files": file_results, "raw_rows_inserted": total_raw}
        except Exception:
            conn.rollback()
            raise
        finally:
            cur.close()


def batch_status(batch_id: int) -> dict[str, Any]:
    with db_conn() as conn:
        cur = conn.cursor(dictionary=True)
        cur.execute(
            "SELECT batch_id,batch_key,batch_label,processing_status,upload_started_at,completed_at,file_count,source_row_count,error_count FROM ingest_batch WHERE batch_id=%s",
            (batch_id,),
        )
        batch = cur.fetchone()
        if not batch:
            raise KeyError(f"Batch {batch_id} not found")
        cur.execute(
            "SELECT file_id,original_filename,source_type,detected_format,processing_status,source_row_count,error_count,duplicate_of_file_id,coverage_type,coverage_start_date,coverage_end_date FROM ingest_file WHERE batch_id=%s ORDER BY file_id",
            (batch_id,),
        )
        files = cur.fetchall()
        cur.execute(
            "SELECT severity,error_code,error_message,file_id,sheet_id,source_row_number,created_at FROM ingest_error WHERE batch_id=%s ORDER BY error_id",
            (batch_id,),
        )
        errors = cur.fetchall()
        cur.close()
        for obj in [batch, *files, *errors]:
            for k, v in list(obj.items()):
                if isinstance(v, (date, datetime)):
                    obj[k] = v.isoformat()
        return {"batch": batch, "files": files, "issues": errors}


def audit_batch(batch_id: int) -> dict[str, Any]:
    s = batch_status(batch_id)
    files = s["files"]
    issues = s["issues"]
    blocking = [x for x in issues if str(x.get("severity", "")).upper() in {"ERROR", "CRITICAL"}]
    checks = {
        "has_files": len(files) > 0,
        "all_files_terminal": all(f["processing_status"] in {"INGESTED", "QA_HOLD"} for f in files),
        "no_blocking_errors": len(blocking) == 0,
        "no_unclassified_files": all(f["source_type"] != "unclassified" for f in files),
        "all_workbooks_have_rows_or_duplicate": all((f["source_row_count"] or 0) > 0 or f["duplicate_of_file_id"] is not None for f in files),
    }
    return {"batch_id": batch_id, "pass": all(checks.values()), "checks": checks, "blocking_issues": blocking}


def recent_batches(limit: int = 20) -> list[dict[str, Any]]:
    limit = max(1, min(int(limit), 100))
    with db_conn() as conn:
        cur = conn.cursor(dictionary=True)
        cur.execute(
            f"SELECT batch_id,batch_label,processing_status,upload_started_at,completed_at,file_count,source_row_count,error_count FROM ingest_batch ORDER BY batch_id DESC LIMIT {limit}"
        )
        rows = cur.fetchall()
        cur.close()
    for row in rows:
        for k, v in list(row.items()):
            if isinstance(v, datetime):
                row[k] = v.isoformat(sep=' ')
            elif isinstance(v, date):
                row[k] = v.isoformat()
    return rows


def release_batch(batch_id: int) -> dict[str, Any]:
    audit = audit_batch(batch_id)
    if not audit["pass"]:
        return {"batch_id": batch_id, "released": False, "audit": audit}
    with db_conn() as conn:
        cur = conn.cursor()
        cur.execute(
            "UPDATE ingest_batch SET processing_status='RAW_RELEASED',completed_at=CURRENT_TIMESTAMP(6) WHERE batch_id=%s AND processing_status IN ('INGESTED_PENDING_QA','QA_READY')",
            (batch_id,),
        )
        changed = cur.rowcount
        conn.commit()
        cur.close()
    return {"batch_id": batch_id, "released": bool(changed), "audit": audit}
